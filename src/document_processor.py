import io
import requests
from pathlib import Path
from PIL import Image
import pytesseract
import PyPDF2
from docx import Document
from docx.shared import Inches
import openpyxl

class DocumentProcessor:
    def __init__(self, logger):
        self.logger = logger

    def process_pdf(self, file_path):
        text = ""
        try:
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    text += page.extract_text()
        except Exception as e:
            self.logger.error(f"PDF处理失败: {e}")
        return text

    def process_image(self, file_path):
        try:
            img = Image.open(file_path)
            text = pytesseract.image_to_string(img, lang='chi_sim+eng')
            return text, img
        except Exception as e:
            self.logger.error(f"图片处理失败: {e}")
            return "", None

    def process_word(self, file_path):
        text = ""
        try:
            doc = Document(file_path)
            for para in doc.paragraphs:
                text += para.text + "\n"
        except Exception as e:
            self.logger.error(f"Word处理失败: {e}")
        return text

    def process_excel(self, file_path):
        text = ""
        try:
            wb = openpyxl.load_workbook(file_path)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    text += " | ".join([str(cell) if cell else "" for cell in row]) + "\n"
        except Exception as e:
            self.logger.error(f"Excel处理失败: {e}")
        return text
